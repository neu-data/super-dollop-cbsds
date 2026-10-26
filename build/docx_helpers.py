"""Shared helpers for the course Word documents (python-docx).

Documents are built on the Neudata document template (build/templates/Neudata-Document.docx,
from https://github.com/neu-data/word-templates). The template supplies the styles
(Title, Subtitle, Heading 1-3 in Century Gothic, Normal in Arial 10.5 pt, Key Findings,
Neudata Small, Caption), the header with the logo and the Neudata footer with page numbers.
"""
import copy
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, "build", "templates", "Neudata-Document.docx")

NAVY = RGBColor(0x04, 0x24, 0x2F)
TEAL = RGBColor(0x05, 0x5F, 0x56)
BLUE = RGBColor(0x0B, 0x37, 0x6C)
GREY = RGBColor(0x8A, 0x99, 0xA3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CODE_FONT = "Consolas"


def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _borders(table, color="C9D6DC", size=4):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(size)); e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)


class Doc:
    def __init__(self, title, subtitle=None, header_text=None, byline=None):
        self.d = Document(TEMPLATE)
        body = self.d.element.body
        for el in list(body):
            if el.tag != qn("w:sectPr"):
                body.remove(el)
        # Header: replace the template's "Document title" text, keep the logo
        hp = self.d.sections[0].header.paragraphs[0]
        for r in hp.runs:
            if r.text.strip() == "Document title":
                r.text = header_text or title
        self.d.core_properties.title = title
        self.d.core_properties.author = "Neudata Consulting Ltd"
        self.d.add_paragraph(title, style="Title")
        if subtitle:
            self.d.add_paragraph(subtitle, style="Subtitle")
        if byline:
            self.d.add_paragraph(byline, style="Neudata Small")

    # ------------------------------------------------------------ content
    def h1(self, t): return self.d.add_heading(t, 1)
    def h2(self, t): return self.d.add_heading(t, 2)
    def h3(self, t): return self.d.add_heading(t, 3)

    def p(self, text, italic=False, size=None, color=None, bold=False):
        para = self.d.add_paragraph()
        self._runs(para, text, italic=italic, size=size, color=color, bold=bold)
        return para

    def _runs(self, para, text, italic=False, size=None, color=None, bold=False):
        """Supports `code` spans inside text."""
        for i, part in enumerate(text.split("`")):
            if not part:
                continue
            r = para.add_run(part)
            r.italic = italic; r.bold = bold
            if size: r.font.size = Pt(size)
            if color: r.font.color.rgb = color
            if i % 2 == 1:
                r.font.name = CODE_FONT
                r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CODE_FONT)
                r.font.size = Pt(size - 1 if size else 9.5)

    def bullets(self, items, style="List Bullet"):
        for it in items:
            para = self.d.add_paragraph(style=style)
            self._runs(para, it)

    def numbered(self, items):
        self.bullets(items, "List Number")

    def code(self, text):
        t = self.d.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.cell(0, 0)
        _shade(c, "EAF2F5")
        c.paragraphs[0].text = ""
        for i, line in enumerate(text.strip("\n").split("\n")):
            para = c.paragraphs[0] if i == 0 else c.add_paragraph()
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            r = para.add_run(line)
            r.font.name = CODE_FONT; r.font.size = Pt(9)
            r.font.color.rgb = NAVY
            r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CODE_FONT)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def answer_box(self, lines=4, label="Your answer"):
        t = self.d.add_table(rows=1, cols=1)
        _borders(t, "5FA39B", 6)
        c = t.cell(0, 0)
        r = c.paragraphs[0].add_run(label); r.font.size = Pt(8); r.font.color.rgb = GREY
        for _ in range(lines):
            c.add_paragraph("")
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def table(self, rows, widths=None, size=9, header=True):
        """Neudata table look: navy header row, light mist striping."""
        t = self.d.add_table(rows=len(rows), cols=len(rows[0]))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t)
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                c = t.cell(i, j)
                c.text = ""
                para = c.paragraphs[0]
                para.paragraph_format.space_after = Pt(1)
                is_head = i == 0 and header
                self._runs(para, str(val), size=size, bold=is_head, color=WHITE if is_head else None)
                if is_head:
                    _shade(c, "04242F")
                elif i % 2 == 0:
                    _shade(c, "EAF2F5")
                if widths:
                    c.width = Cm(widths[j])
        if header:
            _repeat_header(t.rows[0])
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def note(self, text):
        """Key message in the template's 'Key Findings' style (tinted box, teal border)."""
        para = self.d.add_paragraph(style="Key Findings")
        self._runs(para, text)

    def image(self, path, width_cm=16, caption=None):
        if not os.path.exists(path):
            self.p(f"[missing image: {os.path.relpath(path, ROOT)}]", italic=True, color=GREY)
            return
        self.d.add_picture(path, width=Cm(width_cm))
        self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if caption:
            cp = self.d.add_paragraph(caption, style="Caption")
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def save(self, path):
        self.d.save(path)
        print("Saved", path)
