"""Shared layout helpers for the course slide decks (python-pptx).

Slides are built on the Neudata PowerPoint template (build/templates/Neudata-Template.pptx,
from https://github.com/neu-data/powerpoint-template). The template supplies the theme
(Century Gothic, Neudata colours), the logo, the "Neudata | #ClearDataClearImpact"
footer and the layouts: Title Slide, Section Header, Title Only and Closing - Copyright.
Speaker notes carry the detail.
"""
import copy
import os
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn
from PIL import Image

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(ROOT, "figures")
TEMPLATE = os.path.join(ROOT, "build", "templates", "Neudata-Template.pptx")

# Neudata palette (see https://github.com/neu-data/.github/blob/main/BRAND.md)
TEAL = RGBColor(0x0D, 0x73, 0x77)       # PowerPoint title teal
DEEP_TEAL = RGBColor(0x05, 0x5F, 0x56)
NAVY = RGBColor(0x04, 0x24, 0x2F)
BLUE = RGBColor(0x0B, 0x37, 0x6C)
NOTE_BLUE = RGBColor(0x00, 0x70, 0xC0)
MIST = RGBColor(0xEA, 0xF2, 0xF5)
GREY = RGBColor(0x8A, 0x99, 0xA3)
TEXT = RGBColor(0x1F, 0x2D, 0x33)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.47)
TOP = Inches(1.42)          # below the title rule
BOTTOM = Inches(6.72)       # above the logo and footer


class Deck:
    def __init__(self):
        self.prs = Presentation(TEMPLATE)
        # remove the template's example slides, keep its masters and layouts
        sldIdLst = self.prs.slides._sldIdLst
        for sldId in list(sldIdLst):
            self.prs.part.drop_rel(sldId.rId)
            sldIdLst.remove(sldId)
        self.layout = {l.name: l for l in self.prs.slide_layouts}
        self.n = 0

    # ---------------------------------------------------------------- basics
    def _slide(self, layout, notes=None):
        s = self.prs.slides.add_slide(self.layout[layout])
        self.n += 1
        # python-pptx does not copy slide-number placeholders; clone it and add the field
        for lp in self.layout[layout].placeholders:
            if lp.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
                ph = s.shapes.clone_placeholder(lp)
                sp = s.shapes._spTree[-1]
                txBody = sp.find(qn("p:txBody"))
                if txBody is None:
                    txBody = etree.SubElement(sp, qn("p:txBody"))
                    etree.SubElement(txBody, qn("a:bodyPr")); etree.SubElement(txBody, qn("a:lstStyle"))
                for p in txBody.findall(qn("a:p")):
                    txBody.remove(p)
                p = etree.SubElement(txBody, qn("a:p"))
                fld = etree.SubElement(p, qn("a:fld"), id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", type="slidenum")
                etree.SubElement(fld, qn("a:rPr"), lang="en-GB")
                t = etree.SubElement(fld, qn("a:t")); t.text = str(self.n)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def _run_font(self, r, size, color, bold=False, italic=False):
        f = r.font
        f.size, f.bold, f.italic = Pt(size), bold, italic
        f.color.rgb = color

    def _text(self, s, text, x, y, w, h, size=18, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
              anchor=MSO_ANCHOR.TOP, italic=False):
        tb = s.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.05)
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run(); r.text = line
            self._run_font(r, size, color, bold, italic)
        return tb

    def _title(self, s, title, sub=None):
        s.shapes.title.text = title
        if sub:
            self._text(s, sub, MARGIN, TOP - Inches(0.05), W - 2 * MARGIN, Inches(0.45), size=16,
                       color=NOTE_BLUE)

    def _bullets(self, s, items, x, y, w, h, size=20):
        """items: list of str or (str, level). Level 1 is indented, smaller."""
        tb = s.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        first = True
        for it in items:
            text, lvl = (it, 0) if isinstance(it, str) else it
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_after = Pt(10 if lvl == 0 else 5)
            b = p.add_run()
            b.text = ("●  " if lvl == 0 else "        –  ")
            self._run_font(b, (size - 6) if lvl == 0 else size - 3, TEAL if lvl == 0 else GREY)
            r = p.add_run(); r.text = text
            self._run_font(r, size if lvl == 0 else size - 3, TEXT if lvl == 0 else GREY)
        return tb

    def _image(self, s, name, x, y, w, h):
        path = os.path.join(FIG, name)
        if not os.path.exists(path):
            self._text(s, f"[figure {name} not yet built]", x, y, w, h, size=14, color=BLUE)
            return
        iw, ih = Image.open(path).size
        scale = min(w / iw, h / ih)
        dw, dh = int(iw * scale), int(ih * scale)
        s.shapes.add_picture(path, x + (w - dw) // 2, y + (h - dh) // 2, dw, dh)

    def _box(self, s, text, x, y, w, h, size=17):
        """Key-message box: mist fill with a teal left bar (Neudata 'Key Findings' look)."""
        shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        shp.fill.solid(); shp.fill.fore_color.rgb = MIST
        shp.line.fill.background(); shp.shadow.inherit = False
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.09), h)
        bar.fill.solid(); bar.fill.fore_color.rgb = DEEP_TEAL
        bar.line.fill.background(); bar.shadow.inherit = False
        tf = shp.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3); tf.margin_right = Inches(0.2)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = p.add_run(); r.text = line
            self._run_font(r, size, NAVY)
        return shp

    def _table(self, s, rows, x, y, w, col_w=None, size=15, row_h=0.5):
        nr, nc = len(rows), len(rows[0])
        shape = s.shapes.add_table(nr, nc, x, y, w, Inches(row_h * nr))
        t = shape.table
        if col_w:
            tot = sum(col_w)
            for j, cw in enumerate(col_w):
                t.columns[j].width = int(w * cw / tot)
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                c = t.cell(i, j)
                c.text = ""
                r = c.text_frame.paragraphs[0].add_run(); r.text = str(val)
                c.margin_left = c.margin_right = Inches(0.08)
                c.margin_top = c.margin_bottom = Inches(0.04)
                c.fill.solid()
                if i == 0:
                    c.fill.fore_color.rgb = NAVY
                    self._run_font(r, size, WHITE, bold=True)
                else:
                    c.fill.fore_color.rgb = MIST if i % 2 == 0 else WHITE
                    self._run_font(r, size, TEXT)
        return shape

    # ------------------------------------------------------------ slide types
    def title_slide(self, title, subtitle, date_text, notes=None):
        s = self._slide("Title Slide", notes)
        s.shapes.title.text = title
        for ph in s.placeholders:
            t = ph.placeholder_format.type
            if t == PP_PLACEHOLDER.SUBTITLE:
                ph.text = subtitle
            elif t == PP_PLACEHOLDER.DATE:
                ph.text = date_text

    def section(self, number, title, sub="", notes=None):
        s = self._slide("Section Header", notes)
        s.shapes.title.text = title
        for ph in s.placeholders:
            if ph.placeholder_format.type == PP_PLACEHOLDER.BODY:
                ph.text = f"{number}. {sub}" if sub else number

    def closing(self, notes=None):
        self._slide("Closing - Copyright", notes)

    def bullets(self, title, items, notes=None, sub=None, size=22, box=None):
        s = self._slide("Title Only", notes)
        self._title(s, title, sub)
        top = TOP + (Inches(0.5) if sub else Inches(0.1))
        bh = Inches(1.05) if box else 0
        self._bullets(s, items, MARGIN, top, W - 2 * MARGIN, BOTTOM - top - bh, size)
        if box:
            self._box(s, box, MARGIN, BOTTOM - Inches(0.95), W - 2 * MARGIN, Inches(0.9))
        return s

    def image(self, title, img, notes=None, sub=None, caption=None):
        s = self._slide("Title Only", notes)
        self._title(s, title, sub)
        top = TOP + (Inches(0.5) if sub else Inches(0.05))
        ch = Inches(0.5) if caption else 0
        self._image(s, img, MARGIN, top, W - 2 * MARGIN, BOTTOM - top - ch)
        if caption:
            self._text(s, caption, MARGIN, BOTTOM - Inches(0.45), W - 2 * MARGIN, Inches(0.45), size=15,
                       color=NOTE_BLUE, align=PP_ALIGN.CENTER)
        return s

    def text_image(self, title, items, img, notes=None, sub=None, size=19, split=0.42, box=None):
        s = self._slide("Title Only", notes)
        self._title(s, title, sub)
        top = TOP + (Inches(0.5) if sub else Inches(0.1))
        tw = int((W - 2 * MARGIN) * split)
        bh = Inches(1.05) if box else 0
        self._bullets(s, items, MARGIN, top, tw, BOTTOM - top - bh, size)
        self._image(s, img, MARGIN + tw + Inches(0.2), top, W - 2 * MARGIN - tw - Inches(0.2),
                    BOTTOM - top - bh)
        if box:
            self._box(s, box, MARGIN, BOTTOM - Inches(0.95), W - 2 * MARGIN, Inches(0.9), size=16)
        return s

    def table(self, title, rows, notes=None, sub=None, col_w=None, size=15, items_after=None,
              box=None, row_h=0.5):
        s = self._slide("Title Only", notes)
        self._title(s, title, sub)
        top = TOP + (Inches(0.5) if sub else Inches(0.15))
        self._table(s, rows, MARGIN, top, W - 2 * MARGIN, col_w, size, row_h=row_h)
        if items_after:
            y = top + Inches(row_h * len(rows)) + Inches(0.3)
            self._bullets(s, items_after, MARGIN, y, W - 2 * MARGIN, BOTTOM - y, size=19)
        if box:
            self._box(s, box, MARGIN, BOTTOM - Inches(0.95), W - 2 * MARGIN, Inches(0.9), size=16)
        return s

    def two_col(self, title, left, right, notes=None, sub=None, heads=None, size=19, box=None):
        s = self._slide("Title Only", notes)
        self._title(s, title, sub)
        top = TOP + (Inches(0.5) if sub else Inches(0.1))
        cw = (W - 2 * MARGIN - Inches(0.4)) // 2
        bh = Inches(1.05) if box else 0
        for k, items in enumerate([left, right]):
            x = MARGIN + k * (cw + Inches(0.4))
            y = top
            if heads:
                self._text(s, heads[k], x, y, cw, Inches(0.5), size=21, color=BLUE, bold=True)
                y += Inches(0.55)
            self._bullets(s, items, x, y, cw, BOTTOM - y - bh, size)
        if box:
            self._box(s, box, MARGIN, BOTTOM - Inches(0.95), W - 2 * MARGIN, Inches(0.9), size=16)
        return s

    def save(self, path):
        self.prs.save(path)
        print(f"Saved {path} ({self.n} slides)")
